"""Core document operations. All output is written to a separate destination."""
from __future__ import annotations
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata
from pypdf import PdfReader, PdfWriter
from openpyxl import Workbook, load_workbook

@dataclass
class Result:
    source: str
    status: str
    detail: str = ""
    output: str = ""

def pdfs(folder: Path) -> list[Path]:
    return sorted(p for p in Path(folder).rglob('*') if p.is_file() and p.suffix.lower() == '.pdf')

def read_pdf(path: Path) -> str:
    return '\n'.join(page.extract_text() or '' for page in PdfReader(str(path)).pages)

def normalize(value: str) -> str:
    return re.sub(r'\s+', ' ', ''.join(c for c in unicodedata.normalize('NFKD', value.casefold()) if not unicodedata.combining(c))).strip()

def clean_name(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', value).strip(' .')[:140] or 'documento'

def money(value: str) -> Decimal | None:
    value = re.sub(r'[^\d,.]', '', value)
    if not value: return None
    if ',' in value: value = value.replace('.', '').replace(',', '.')
    elif value.count('.') > 1: value = value.replace('.', '')
    try: return Decimal(value).quantize(Decimal('0.01'))
    except InvalidOperation: return None

def valid_cnpj(value: str) -> bool:
    """Valida CNPJ numérico pelo módulo 11."""
    digits = re.sub(r'\D', '', value)
    if len(digits) != 14 or len(set(digits)) == 1:
        return False
    for size, weights in ((12, (5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2)),
                          (13, (6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2))):
        remainder = sum(int(digits[i]) * weights[i] for i in range(size)) % 11
        if int(digits[size]) != (0 if remainder < 2 else 11 - remainder):
            return False
    return True

def extract_fields(path: Path) -> dict:
    text = read_pdf(path)
    patterns = {
        'cnpj': r'\b\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}\b|\b\d{14}\b',
        'numero': r'(?:N[º°o.]?\s*(?:da\s*)?(?:nota|nf)|nota\s*fiscal\s*(?:n[º°o.]?)?|documento\s*n[º°o.]?)\s*[:#-]?\s*(\d{1,12})',
        'valor': r'(?:valor\s*(?:total|cobrado|pago)?|total\s*(?:da\s*nota)?)\s*[:R$\s]*([\d.]+,\d{2}|\d+\.\d{2})',
    }
    fields = {}
    for key, pattern in patterns.items():
        match = next((m for m in re.finditer(pattern, text, re.I) if key != 'cnpj' or valid_cnpj(m.group(0))), None)
        fields[key] = (match.group(1) if key != 'cnpj' else match.group(0)) if match else None
    if fields['valor'] is not None: fields['valor'] = money(fields['valor'])
    return fields

def merge_pair(receipt: Path, invoice: Path, destination: Path) -> Path:
    receipt, invoice, destination = map(Path, (receipt, invoice, destination))
    if destination.resolve() in (receipt.resolve(), invoice.resolve()):
        raise ValueError('O arquivo de saída não pode substituir uma entrada.')
    writer = PdfWriter()
    for file in (receipt, invoice):
        for page in PdfReader(str(file)).pages: writer.add_page(page)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open('wb') as output: writer.write(output)
    return destination

def merge_receipt_to_many(receipt: Path, invoice_files: list[Path], output_dir: Path) -> list[Result]:
    output_dir = Path(output_dir)
    results = []
    for invoice in invoice_files:
        try:
            output = output_dir / f'combinado_{clean_name(invoice.stem)}.pdf'
            if output.exists(): raise FileExistsError(f'Saída já existe: {output.name}')
            merge_pair(receipt, invoice, output)
            results.append(Result(invoice.name, 'OK', 'Comprovante antes da nota', str(output)))
        except Exception as exc: results.append(Result(invoice.name, 'ERRO', str(exc)))
    return results

def search_pdfs(folder: Path, term: str, year: int | None = None, content: bool = False) -> list[Result]:
    results = []
    for file in pdfs(folder):
        try:
            haystack = file.name + ('\n' + read_pdf(file) if content else '')
            if normalize(term) in normalize(haystack) and (year is None or str(year) in normalize(haystack + ' ' + str(file.parent))):
                results.append(Result(file.name, 'ENCONTRADO', str(file)))
        except Exception as exc: results.append(Result(file.name, 'ERRO', str(exc)))
    return results

def rename_suggestions(files: list[Path]) -> list[Result]:
    results = []
    for file in files:
        try:
            fields = extract_fields(file)
            number = fields['numero'] or 'SEM_NUMERO'
            amount = str(fields['valor']).replace('.', ',') if fields['valor'] is not None else 'SEM_VALOR'
            name = clean_name(f'{file.stem} NF_{number} R$_{amount}') + '.pdf'
            results.append(Result(file.name, 'REVISAR' if 'SEM_' in name else 'SUGESTAO', name))
        except Exception as exc: results.append(Result(file.name, 'ERRO', str(exc)))
    return results

def export_report(results: list[Result], destination: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = 'Resultados'
    ws.append(['Arquivo', 'Status', 'Detalhes', 'Saída'])
    for item in results: ws.append([item.source, item.status, item.detail, item.output])
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for column, width in {'A':36,'B':18,'C':65,'D':65}.items(): ws.column_dimensions[column].width = width
    destination = Path(destination); destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination); return destination

def match_by_spreadsheet(excel: Path, invoices: list[Path], receipts: list[Path], output_dir: Path, invoice_column: str = 'NF') -> list[Result]:
    """Match invoice numbers in Excel against filenames; ambiguous matches require review."""
    wb = load_workbook(excel, read_only=True, data_only=True)
    try:
        ws = wb.active; rows = ws.iter_rows(values_only=True)
        first = next(rows, None)
        if first is None or not any(v is not None and str(v).strip() for v in first):
            raise ValueError('A planilha está vazia ou não possui cabeçalho.')
        headers = [str(v or '').strip().casefold() for v in first]
        try: idx = headers.index(invoice_column.strip().casefold())
        except ValueError: raise ValueError(f'Coluna {invoice_column!r} não encontrada na planilha')
        results = []
        for row in rows:
            if idx >= len(row) or row[idx] is None: continue
            number = str(row[idx]).strip()
            if number.endswith('.0'): number = number[:-2]
            pattern = re.compile(r'(?<!\d)' + re.escape(number) + r'(?!\d)')
            notes = [p for p in invoices if pattern.search(p.stem)]
            proofs = [p for p in receipts if pattern.search(p.stem)]
            if len(notes) != 1 or len(proofs) != 1:
                results.append(Result(number, 'REVISAR', f'{len(notes)} notas e {len(proofs)} comprovantes encontrados'))
                continue
            try:
                dest = Path(output_dir) / f'NF_{clean_name(number)}.pdf'
                if dest.exists(): raise FileExistsError('Arquivo de saída já existe')
                merge_pair(proofs[0], notes[0], dest)
                results.append(Result(number, 'OK', 'Comprovante antes da nota', str(dest)))
            except Exception as exc: results.append(Result(number, 'ERRO', str(exc)))
        return results
    finally:
        wb.close()
